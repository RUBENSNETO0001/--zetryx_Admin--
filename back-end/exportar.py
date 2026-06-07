from io import BytesIO
from flask import Blueprint, jsonify, send_file
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import date, datetime
from decimal import Decimal

exportar_bp = Blueprint("exportar", __name__)

SECOES = [
    ("DADOS PESSOAIS", "1F4E79", [
        ("Matrícula",           "matricula"),
        ("Nome Completo",       "nome_completo"),
        ("CPF",                 "cpf"),
        ("Data de Nascimento",  "data_nascimento"),
        ("E-mail",              "email"),
        ("Telefone",            "telefone"),
        ("Estado Civil",        "estado_civil"),
        ("Modalidade Auxílio",  "modalidade_auxilio"),
        ("Status",              "status"),
        ("Data de Cadastro",    "data_cadastro"),
    ]),
    ("ENDEREÇO", "2E75B6", [
        ("Rua",                 "rua"),
        ("Número",              "numero"),
        ("Bairro",              "bairro"),
        ("Cidade",              "cidade"),
        ("CEP",                 "cep"),
        ("Complemento",         "complemento"),
    ]),
    ("CURSO / MATRÍCULA", "375623", [
        ("Curso",               "nome_curso"),
        ("Tipo de Curso",       "tipo_curso"),
        ("Data de Matrícula",   "data_matricula"),
        ("Qtd. Disciplinas",    "qtd_disciplinas_matriculadas"),
    ]),
    ("DADOS SOCIOECONÔMICOS", "833C00", [
        ("Tipo de Moradia",         "tipo_moradia"),
        ("Mora em Zona",            "mora_em"),
        ("Renda Bruta Familiar",    "renda_bruta_total_familiar"),
        ("Qtd. Pessoas na Casa",    "quantidade_pessoas_casa"),
    ]),
    ("BENEFÍCIOS", "7030A0", [
        ("Recebe Auxílio Bolsa?",       "recebe_auxilio_bolsa"),
        ("Nome do Benefício",           "nome_beneficio"),
        ("Bolsa Família / CadÚnico?",   "beneficiario_bolsa_familia_cadunico"),
        ("Recebe BPC?",                 "recebe_bpc"),
    ]),
    ("DECLARAÇÕES", "C00000", [
        ("Lei de Cotas?",       "lei_cotas"),
        ("Possui Deficiência?", "possui_deficiencia"),
        ("Quilombola?",         "origem_quilombola"),
        ("Estrangeiro?",        "estrangeiro"),
    ]),
    ("PERFIL / REQUISITOS", "4472C4", [
        ("Escolaridade (EM)",       "escolaridade"),
        ("Renda Per Capita (R$)",   "renda_per_capita"),
    ]),
    ("DADOS BANCÁRIOS", "006400", [
        ("Banco",               "nome_instituicao"),
        ("Possui Conta?",       "possui_conta"),
        ("Tipo de Conta",       "tipo_conta"),
        ("Número da Conta",     "numero_conta"),
        ("Agência",             "agencia"),
        ("Variação Poupança",   "variacao_poupanca"),
    ]),
    ("DOCUMENTOS", "595959", [
        ("Doc 1 – Título",      "doc1_titulo"),
        ("Doc 1 – Status",      "doc1_status"),
        ("Doc 1 – Rejeição",    "doc1_rejeicao"),
        ("Doc 2 – Título",      "doc2_titulo"),
        ("Doc 2 – Status",      "doc2_status"),
        ("Doc 2 – Rejeição",    "doc2_rejeicao"),
    ]),
    ("PONTUAÇÃO", "1F4E79", [
        ("Pts – Renda",                 "pts_renda"),
        ("Pts – Escolaridade",          "pts_escolaridade"),
        ("Pts – Lei de Cotas",          "pts_lei_cotas"),
        ("Pts – Deficiência (part.)",   "pts_deficiencia_participante"),
        ("Pts – Quilombola",            "pts_quilombola"),
        ("Pts – Estrangeiro",           "pts_estrangeiro"),
        ("Pts – Aluguel",               "pts_aluguel"),
        ("Pts – Zona Rural",            "pts_zona_rural"),
        ("Pts – Deficiência (fam.)",    "pts_deficiencia_familia"),
        ("Pts – Doença Crônica (fam.)", "pts_doenca_cronica_familia"),
        ("Pts – Benefício Social",      "pts_beneficio_social"),
        ("Ajuste Manual",               "ajuste_manual"),
        ("Observação Ajuste",           "observacao_ajuste"),
        ("Pontuação Total",             "pontuacao_total"),
        ("Classificação",               "classificacao"),
    ]),
]

BOOL_KEYS = {
    "recebe_auxilio_bolsa", "beneficiario_bolsa_familia_cadunico", "recebe_bpc",
    "lei_cotas", "possui_deficiencia", "origem_quilombola", "estrangeiro", "possui_conta",
}

STATUS_CORES = {"Novo": "4472C4", "Revisando": "ED7D31", "Finalizado": "70AD47"}
DOC_CORES    = {"aprovado": "70AD47", "recusado": "FF0000", "pendente": "FFC000"}
CLASS_CORES  = {
    "Vulnerabilidade I":   ("C00000", "FFFFFF", True),
    "Vulnerabilidade II":  ("ED7D31", "FFFFFF", True),
    "Vulnerabilidade III": ("FFC000", "000000", False),
    "Não prioritário":     ("D9D9D9", "000000", False),
}

_thin = Side(style="thin", color="BFBFBF")
_brd  = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)


def _fmt(p, key):
    v = p.get(key)
    if v is None:
        return ""
    if key in BOOL_KEYS:
        if isinstance(v, bool): return "Sim" if v else "Não"
        if isinstance(v, int) and v in (0, 1): return "Sim" if v else "Não"
    if isinstance(v, (date, datetime)):
        return str(v)
    if isinstance(v, Decimal):
        return float(v)
    if isinstance(v, bool):
        return "Sim" if v else "Não"
    return v


def _cel(ws, row, col, value="", bold=False, bg="FFFFFF",
         fg="000000", size=9, halign="center", wrap=False):
    c = ws.cell(row=row, column=col, value=value)
    c.font      = Font(name="Arial", bold=bold, color=fg, size=size)
    c.fill      = PatternFill("solid", start_color=bg)
    c.alignment = Alignment(horizontal=halign, vertical="center", wrap_text=wrap)
    c.border    = _brd
    return c


def gerar_xlsx_todos(participantes: list) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Participantes"
    ws.sheet_view.showGridLines = False

    total_cols = sum(len(campos) for _, _, campos in SECOES)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_cols)
    _cel(ws, 1, 1,
         "SISTEMA ZETRYX – RELATÓRIO DE PARTICIPANTES",
         bold=True, bg="1F4E79", fg="FFFFFF", size=13)
    ws.row_dimensions[1].height = 26

    col = 1
    for nome_sec, cor, campos in SECOES:
        ncols = len(campos)
        if ncols > 1:
            ws.merge_cells(start_row=2, start_column=col,
                           end_row=2, end_column=col + ncols - 1)
        _cel(ws, 2, col, nome_sec, bold=True, bg=cor, fg="FFFFFF", size=9)
        ws.row_dimensions[2].height = 18

        for i, (label, _) in enumerate(campos):
            c3 = ws.cell(row=3, column=col + i, value=label)
            c3.font      = Font(name="Arial", bold=True, color="FFFFFF", size=8)
            c3.fill      = PatternFill("solid", start_color=cor)
            c3.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c3.border    = _brd
            ws.column_dimensions[get_column_letter(col + i)].width = 18
        col += ncols

    ws.row_dimensions[3].height = 32

    BG_PAR = "F2F7FC"
    BG_IMP = "FFFFFF"

    for ri, p in enumerate(participantes):
        rn  = 4 + ri
        bg  = BG_PAR if ri % 2 == 0 else BG_IMP
        col = 1
        for _, _, campos in SECOES:
            for label, key in campos:
                v  = _fmt(p, key)
                cv = ws.cell(row=rn, column=col, value=v)
                cv.font      = Font(name="Arial", size=9)
                cv.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cv.border    = _brd

                if key == "status" and v in STATUS_CORES:
                    cv.fill = PatternFill("solid", start_color=STATUS_CORES[v])
                    cv.font = Font(name="Arial", size=9, bold=True, color="FFFFFF")
                elif key in ("doc1_status", "doc2_status") and v in DOC_CORES:
                    cv.fill = PatternFill("solid", start_color=DOC_CORES[v])
                    cv.font = Font(name="Arial", size=9, bold=True, color="FFFFFF")
                elif key == "classificacao" and v in CLASS_CORES:
                    cbg, cfg, cbold = CLASS_CORES[v]
                    cv.fill = PatternFill("solid", start_color=cbg)
                    cv.font = Font(name="Arial", size=9, bold=cbold, color=cfg)
                else:
                    cv.fill = PatternFill("solid", start_color=bg)
                col += 1
        ws.row_dimensions[rn].height = 16

    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:{get_column_letter(total_cols)}3"

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf

@exportar_bp.route("/api/participantes/exportar", methods=["GET"])
def exportar_todos():
    try:
        from app import get_db
    except ImportError:
        return jsonify({"error": "Não foi possível importar app.py"}), 500

    try:
        conn   = get_db()
        cursor = conn.cursor(dictionary=True)
        cursor.execute("""
            SELECT
                p.id_participante,
                p.matricula, p.nome_completo, p.cpf, p.data_nascimento,
                p.email, p.telefone, p.estado_civil,
                p.modalidade_auxilio, p.status, p.data_cadastro,

                e.rua, e.numero, e.bairro, e.cidade, e.cep, e.complemento,

                c.nome_curso, c.tipo_curso,
                m.data_matricula, m.qtd_disciplinas_matriculadas,

                ds.tipo_moradia, ds.mora_em,
                ds.renda_bruta_total_familiar, ds.quantidade_pessoas_casa,

                pb.recebe_auxilio_bolsa, pb.nome_beneficio,
                pb.beneficiario_bolsa_familia_cadunico, pb.recebe_bpc,

                pd2.lei_cotas, pd2.possui_deficiencia,
                pd2.origem_quilombola, pd2.estrangeiro,

                pr.escolaridade, pr.renda_per_capita,

                db2.possui_conta, db2.tipo_conta,
                db2.numero_conta, db2.agencia, db2.variacao_poupanca,
                bo.nome_instituicao,

                pp.pts_renda, pp.pts_escolaridade, pp.pts_lei_cotas,
                pp.pts_deficiencia_participante, pp.pts_quilombola,
                pp.pts_estrangeiro, pp.pts_aluguel, pp.pts_zona_rural,
                pp.pts_deficiencia_familia, pp.pts_doenca_cronica_familia,
                pp.pts_beneficio_social, pp.ajuste_manual,
                pp.observacao_ajuste, pp.pontuacao_total, pp.classificacao

            FROM Participante p
            LEFT JOIN Endereco_participante e  ON e.id_participante  = p.id_participante
            LEFT JOIN Curso c                  ON c.id_participante  = p.id_participante
            LEFT JOIN Matricula m              ON m.id_participante  = p.id_participante
            LEFT JOIN Dados_Socioeconomicos ds ON ds.id_participante = p.id_participante
            LEFT JOIN Participante_Beneficios pb ON pb.id_socio      = ds.id_socio
            LEFT JOIN perfil_declaracao pd2    ON pd2.id_participante= p.id_participante
            LEFT JOIN Perfil_requisitos pr     ON pr.id_participante = p.id_participante
            LEFT JOIN Dados_Bancarios db2      ON db2.id_participante= p.id_participante
            LEFT JOIN Banco_Oficiais bo        ON bo.id_bancoOf      = db2.id_bancoOf
            LEFT JOIN Pontuacao_Participante pp ON pp.id_participante= p.id_participante
            ORDER BY p.data_cadastro DESC
        """)
        rows = cursor.fetchall()

        # Documentos: pega até 2 por participante
        cursor.execute("""
            SELECT id_participante, titulo_do_pdf, status_validacao, justificativa_rejeicao
            FROM Pdf_Participante
            ORDER BY id_participante, id_pdf
        """)
        docs_raw = cursor.fetchall()
        docs_map = {}
        for d in docs_raw:
            pid = d["id_participante"]
            docs_map.setdefault(pid, []).append(d)

        for row in rows:
            pid  = row["id_participante"]
            docs = docs_map.get(pid, [])
            row["doc1_titulo"]   = docs[0]["titulo_do_pdf"]         if len(docs) > 0 else ""
            row["doc1_status"]   = docs[0]["status_validacao"]      if len(docs) > 0 else ""
            row["doc1_rejeicao"] = docs[0]["justificativa_rejeicao"] or "" if len(docs) > 0 else ""
            row["doc2_titulo"]   = docs[1]["titulo_do_pdf"]         if len(docs) > 1 else ""
            row["doc2_status"]   = docs[1]["status_validacao"]      if len(docs) > 1 else ""
            row["doc2_rejeicao"] = docs[1]["justificativa_rejeicao"] or "" if len(docs) > 1 else ""

        buf = gerar_xlsx_todos(rows)
        return send_file(
            buf,
            as_attachment=True,
            download_name="participantes.xlsx",
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    except Exception as e:
        return jsonify({"error": str(e)}), 500
    finally:
        cursor.close()
        conn.close()